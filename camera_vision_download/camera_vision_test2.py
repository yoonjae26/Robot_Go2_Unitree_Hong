import sys
import time
import base64
import tempfile
import os
import re

import cv2
import numpy as np
import pygame

from openpyxl import load_workbook
from gtts import gTTS
from openai import OpenAI
from insightface.app import FaceAnalysis

from unitree_sdk2py.core.channel import ChannelFactoryInitialize
from unitree_sdk2py.go2.video.video_client import VideoClient


# =========================================================
# 기본 설정
# =========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_FILE = os.path.join(
    BASE_DIR,
    "vision_people.xlsx"
)

FACES_DIR = os.path.join(
    BASE_DIR,
    "faces"
)


# InsightFace cosine similarity 기준
# 높을수록 같은 사람일 가능성이 큼
FACE_SIMILARITY_THRESHOLD = 0.35


# =========================================================
# 한국어 TTS
# 노트북 스피커 출력
# =========================================================
def speak_text(text):

    print("\n[TTS] 한국어 음성을 생성합니다...")

    speech_file = None

    try:

        with tempfile.NamedTemporaryFile(
            suffix=".mp3",
            delete=False
        ) as temp_file:

            speech_file = temp_file.name


        tts = gTTS(
            text=text,
            lang="ko"
        )

        tts.save(speech_file)

        print("[TTS] 음성 생성 완료")
        print("[TTS] 노트북 스피커로 재생합니다...")


        if not pygame.mixer.get_init():

            pygame.mixer.init()


        pygame.mixer.music.load(
            speech_file
        )

        pygame.mixer.music.play()


        while pygame.mixer.music.get_busy():

            pygame.time.Clock().tick(10)


        pygame.mixer.music.unload()

        print("[TTS] 음성 출력 완료")


    except Exception as e:

        print("\n[TTS 오류]")
        print(e)


    finally:

        if speech_file is not None:

            try:

                if os.path.exists(speech_file):

                    os.remove(speech_file)

            except Exception:

                pass


# =========================================================
# 문자열 정리
# =========================================================
def clean_text(value):

    if value is None:

        return ""

    return str(value).strip()


# =========================================================
# 사진 파일명 분리
#
# 한 장:
# gyuyoung.jpg
#
# 여러 장:
# gyuyoung01.jpg, gyuyoung02.jpg
# =========================================================
def split_photo_names(photo_text):

    photo_text = clean_text(
        photo_text
    )


    if not photo_text:

        return []


    photo_names = re.split(
        r"[,;]",
        photo_text
    )


    return [

        name.strip()

        for name in photo_names

        if name.strip()

    ]


# =========================================================
# Cosine Similarity
# =========================================================
def cosine_similarity(
    vector_a,
    vector_b
):

    norm_a = np.linalg.norm(
        vector_a
    )

    norm_b = np.linalg.norm(
        vector_b
    )


    if norm_a == 0 or norm_b == 0:

        return -1.0


    return float(

        np.dot(
            vector_a,
            vector_b
        )

        /

        (
            norm_a
            *
            norm_b
        )

    )


# =========================================================
# InsightFace 초기화
# =========================================================
def initialize_face_model():

    print("")
    print("========================================")
    print("[InsightFace] 얼굴 인식 모델 초기화")
    print("========================================")


    face_model = FaceAnalysis(

        name="buffalo_l",

        providers=[
            "CPUExecutionProvider"
        ]

    )


    # CPU 사용
    # 멀리 있는 작은 얼굴 검출을 위해 960x960 사용
    face_model.prepare(

        ctx_id=-1,

        det_size=(
            960,
            960
        )

    )


    print(
        "[InsightFace] 모델 초기화 완료"
    )


    return face_model


# =========================================================
# 등록 사진에서 가장 큰 얼굴 선택
# =========================================================
def get_largest_face(
    faces
):

    if len(faces) == 0:

        return None


    largest_face = None

    largest_area = 0


    for face in faces:

        bbox = face.bbox


        width = (
            bbox[2]
            -
            bbox[0]
        )

        height = (
            bbox[3]
            -
            bbox[1]
        )


        area = (
            width
            *
            height
        )


        if area > largest_area:

            largest_area = area

            largest_face = face


    return largest_face


# =========================================================
# 엑셀 + 등록 얼굴 DB 읽기
#
# 엑셀 헤더:
#
# name | title | photo | introduction
# =========================================================
def load_people_database(
    face_model
):

    print("")
    print("========================================")
    print("[얼굴 DB] 등록 인물 정보를 읽습니다.")
    print("========================================")


    if not os.path.exists(
        EXCEL_FILE
    ):

        print(
            "[오류] 엑셀 파일을 찾을 수 없습니다."
        )

        print(
            EXCEL_FILE
        )

        return []


    if not os.path.exists(
        FACES_DIR
    ):

        print(
            "[오류] faces 폴더를 찾을 수 없습니다."
        )

        print(
            FACES_DIR
        )

        return []


    try:

        workbook = load_workbook(

            EXCEL_FILE,

            data_only=True

        )


        sheet = workbook.active


        # -------------------------------------------------
        # 헤더 읽기
        # -------------------------------------------------
        headers = {}


        for column_index, cell in enumerate(

            sheet[1],

            start=1

        ):

            header_name = clean_text(
                cell.value
            ).lower()


            if header_name:

                headers[
                    header_name
                ] = column_index


        required_headers = [

            "name",

            "title",

            "photo",

            "introduction"

        ]


        missing_headers = [

            header

            for header in required_headers

            if header not in headers

        ]


        if missing_headers:

            print(
                "[오류] 엑셀 헤더가 올바르지 않습니다."
            )

            print(
                "필요한 헤더:",
                required_headers
            )

            print(
                "없는 헤더:",
                missing_headers
            )

            return []


        people = []


        # -------------------------------------------------
        # 각 사람 읽기
        # -------------------------------------------------
        for row_index in range(

            2,

            sheet.max_row + 1

        ):

            name = clean_text(

                sheet.cell(

                    row=row_index,

                    column=headers[
                        "name"
                    ]

                ).value

            )


            title = clean_text(

                sheet.cell(

                    row=row_index,

                    column=headers[
                        "title"
                    ]

                ).value

            )


            photo_text = clean_text(

                sheet.cell(

                    row=row_index,

                    column=headers[
                        "photo"
                    ]

                ).value

            )


            introduction = clean_text(

                sheet.cell(

                    row=row_index,

                    column=headers[
                        "introduction"
                    ]

                ).value

            )


            if not name:

                continue


            photo_names = split_photo_names(

                photo_text

            )


            if not photo_names:

                print(
                    f"[경고] {name}: 등록 사진이 없습니다."
                )

                continue


            embeddings = []


            # ---------------------------------------------
            # 등록 사진 처리
            # ---------------------------------------------
            for photo_name in photo_names:

                photo_path = os.path.join(

                    FACES_DIR,

                    photo_name

                )


                if not os.path.exists(
                    photo_path
                ):

                    print(
                        f"[경고] 사진 파일이 없습니다: "
                        f"{photo_path}"
                    )

                    continue


                image = cv2.imread(

                    photo_path

                )


                if image is None:

                    print(
                        f"[경고] 사진을 읽을 수 없습니다: "
                        f"{photo_path}"
                    )

                    continue


                faces = face_model.get(

                    image

                )


                if len(faces) == 0:

                    print(
                        f"[경고] {name}: "
                        f"{photo_name}에서 얼굴을 찾지 못했습니다."
                    )

                    continue


                face = get_largest_face(

                    faces

                )


                if face is None:

                    continue


                embedding = (

                    face.normed_embedding

                )


                embeddings.append(

                    embedding

                )


                print(
                    f"[등록 완료] "
                    f"{name} {title} - "
                    f"{photo_name}"
                )


            if len(
                embeddings
            ) == 0:

                print(
                    f"[경고] {name}: "
                    "유효한 얼굴 임베딩이 없습니다."
                )

                continue


            people.append(

                {

                    "name":
                        name,

                    "title":
                        title,

                    "introduction":
                        introduction,

                    "embeddings":
                        embeddings

                }

            )


        print("")

        print(
            f"[얼굴 DB] 총 "
            f"{len(people)}명 등록 완료"
        )


        return people


    except Exception as e:

        print("")
        print(
            "[얼굴 DB 로딩 오류]"
        )
        print(e)

        return []


# =========================================================
# 현재 화면 얼굴 인식
# =========================================================
def recognize_faces(

    frame,

    people,

    face_model

):

    print("")
    print(
        "[InsightFace] 화면의 얼굴을 분석합니다..."
    )


    try:

        faces = face_model.get(

            frame

        )

    except Exception as e:

        print(
            "[InsightFace 오류]"
        )

        print(e)

        return []


    if len(
        faces
    ) == 0:

        print(
            "[InsightFace] 얼굴을 발견하지 못했습니다."
        )

        return []


    print(
        f"[InsightFace] "
        f"{len(faces)}개의 얼굴을 발견했습니다."
    )


    results = []


    # -----------------------------------------------------
    # 각 얼굴마다 등록자 비교
    # -----------------------------------------------------
    for face_index, face in enumerate(

        faces,

        start=1

    ):

        current_embedding = (

            face.normed_embedding

        )


        best_person = None

        best_similarity = -1.0


        # -------------------------------------------------
        # 모든 등록 인물과 비교
        # -------------------------------------------------
        for person in people:

            for registered_embedding in (

                person[
                    "embeddings"
                ]

            ):

                similarity = cosine_similarity(

                    current_embedding,

                    registered_embedding

                )


                if similarity > best_similarity:

                    best_similarity = (

                        similarity

                    )

                    best_person = person


        print(
            f"[얼굴 {face_index}] "
            f"최고 유사도: "
            f"{best_similarity:.3f}"
        )


        # -------------------------------------------------
        # 등록 인물 판정
        # -------------------------------------------------
        if (

            best_person is not None

            and

            best_similarity
            >=
            FACE_SIMILARITY_THRESHOLD

        ):

            print(
                f"[얼굴 {face_index}] "
                f"등록 인물: "
                f"{best_person['name']} "
                f"{best_person['title']}"
            )


            results.append(

                {

                    "recognized":
                        True,

                    "name":
                        best_person[
                            "name"
                        ],

                    "title":
                        best_person[
                            "title"
                        ],

                    "introduction":
                        best_person[
                            "introduction"
                        ],

                    "similarity":
                        best_similarity

                }

            )


        else:

            print(
                f"[얼굴 {face_index}] "
                "등록되지 않은 사람입니다."
            )


            results.append(

                {

                    "recognized":
                        False,

                    "similarity":
                        best_similarity

                }

            )


    return results


# =========================================================
# GPT용 사람 정보 생성
# =========================================================
def make_people_context(

    face_results

):

    if len(
        face_results
    ) == 0:

        return (
            "InsightFace 얼굴 인식 모듈에서는 "
            "현재 화면에서 사람 얼굴을 확인하지 못했습니다."
        )


    context_lines = []


    for index, result in enumerate(

        face_results,

        start=1

    ):

        if result[
            "recognized"
        ]:

            context_lines.append(

                f"{index}번째 얼굴은 "
                f"등록된 사람입니다. "

                f"이름은 "
                f"{result['name']}이고, "

                f"직함은 "
                f"{result['title']}입니다. "

                f"등록된 소개 정보는 다음과 같습니다. "

                f"{result['introduction']}"

            )


        else:

            context_lines.append(

                f"{index}번째 얼굴은 "
                "등록된 인물 데이터베이스와 "
                "일치하지 않는 사람입니다."

            )


    return "\n".join(

        context_lines

    )


# =========================================================
# 현재 프레임 분석
#
# 1. InsightFace 얼굴 식별
# 2. GPT Vision 전체 장면 분석
# 3. 인물 소개 생성
# 4. TTS
# =========================================================
def analyze_frame(

    frame,

    openai_client,

    people,

    face_model

):

    print("")
    print(
        "[AI Vision] 현재 장면을 분석합니다..."
    )


    # =====================================================
    # 1. InsightFace 얼굴 식별
    # =====================================================
    face_results = recognize_faces(

        frame,

        people,

        face_model

    )


    people_context = make_people_context(

        face_results

    )


    # =====================================================
    # 2. 이미지 JPEG 인코딩
    # =====================================================
    success, buffer = cv2.imencode(

        ".jpg",

        frame

    )


    if not success:

        print(
            "[오류] 이미지를 JPEG로 "
            "변환하지 못했습니다."
        )

        return


    image_base64 = base64.b64encode(

        buffer

    ).decode(

        "utf-8"

    )


    # =====================================================
    # 3. GPT Vision
    # =====================================================
    try:

        prompt = f"""
이 이미지는 Unitree Go2 로봇의 전면 카메라 영상입니다.

현재 장면을 자연스러운 한국어 구어체로 설명하세요.

아래에는 별도의 InsightFace 얼굴 인식 시스템이
현재 화면에서 분석한 등록 인물 정보가 있습니다.

사람의 이름과 신원은 반드시 아래의 등록 인물 판정 결과만 사용하세요.
이미지만 보고 사람의 이름이나 신원을 추측하지 마세요.

[등록 인물 판정 결과]

{people_context}


응답 규칙:

1. 이미지에서 실제로 명확하게 보이는 주요 물체와 장면을 짧게 설명하세요.

2. 물체의 위치를 설명할 때는 이미지 화면 기준으로
왼쪽, 중앙, 오른쪽을 판단하세요.

3. 위치가 확실하지 않으면 위치를 말하지 마세요.

4. 거리나 이동 가능 여부 또는 이동 경로는 추측하지 마세요.

5. 등록 인물이 있는 경우
반드시 이름과 직함을 함께 사용하여 소개하세요.

예:
"이 분은 김수연 수석팀장님입니다."

6. 등록된 소개글을 그대로 읽지 말고
사람에게 말하듯 자연스럽게 한두 문장으로 다듬어 소개하세요.

7. 등록되지 않은 사람이 있는 경우에는
이름을 추측하지 말고
"제가 아는 분은 아닙니다"
또는 이에 준하는 자연스러운 표현을 사용하세요.

8. 장면 설명과 인물 소개를
하나의 자연스러운 설명처럼 이어서 말하세요.

9. 너무 길게 말하지 마세요.

10. 번호나 목록 형식으로 답하지 마세요.
"""


        response = openai_client.responses.create(

            model="gpt-4o-mini",

            input=[

                {

                    "role":
                        "user",

                    "content": [

                        {

                            "type":
                                "input_text",

                            "text":
                                prompt

                        },

                        {

                            "type":
                                "input_image",

                            "image_url":
                                "data:image/jpeg;base64,"
                                +
                                image_base64

                        }

                    ]

                }

            ]

        )


        result = response.output_text


        # =================================================
        # 터미널 출력
        # =================================================
        print("")
        print(
            "========================================"
        )

        print(
            "[AI Vision 최종 분석 결과]"
        )

        print(
            result
        )

        print(
            "========================================"
        )


        # =================================================
        # TTS
        # =================================================
        speak_text(

            result

        )


    except Exception as e:

        print("")
        print(
            "[OpenAI Vision API 오류]"
        )

        print(
            e
        )


# =========================================================
# MAIN
# =========================================================
def main():

    # -----------------------------------------------------
    # 1. Unitree 네트워크 초기화
    # -----------------------------------------------------
    if len(
        sys.argv
    ) > 1:

        ChannelFactoryInitialize(

            0,

            sys.argv[
                1
            ]

        )

    else:

        ChannelFactoryInitialize(

            0

        )


    # -----------------------------------------------------
    # 2. OpenAI Client
    # -----------------------------------------------------
    openai_client = OpenAI()


    # -----------------------------------------------------
    # 3. InsightFace 모델 초기화
    # 이미 다운로드된 buffalo_l 모델 사용
    # -----------------------------------------------------
    face_model = initialize_face_model()


    # -----------------------------------------------------
    # 4. 등록 인물 DB 로드
    # -----------------------------------------------------
    people = load_people_database(

        face_model

    )


    if len(
        people
    ) == 0:

        print("")
        print(
            "[주의] 등록된 얼굴이 없습니다."
        )

        print(
            "장면 분석은 가능하지만 "
            "등록 인물 식별은 동작하지 않습니다."
        )


    # -----------------------------------------------------
    # 5. Go2 카메라 초기화
    # -----------------------------------------------------
    print("")
    print(
        "[1] Go2 전면 카메라에 연결합니다."
    )


    video_client = VideoClient()


    video_client.SetTimeout(

        3.0

    )


    video_client.Init()


    # -----------------------------------------------------
    # 6. 화면 생성
    # -----------------------------------------------------
    window_name = (

        "Go2 AI Vision + InsightFace"

    )


    cv2.namedWindow(

        window_name,

        cv2.WINDOW_NORMAL

    )


    print("")
    print(
        "[2] Go2 AI Vision 시작"
    )

    print("")

    print(
        "    a : 현재 장면 분석 + 얼굴 식별 + 음성 출력"
    )

    print(
        "    q : 종료"
    )

    print(
        "    ESC : 종료"
    )

    print("")


    # -----------------------------------------------------
    # 7. 카메라 루프
    # -----------------------------------------------------
    try:

        while True:

            # ---------------------------------------------
            # Go2 카메라 데이터 수신
            # ---------------------------------------------
            try:

                code, data = (

                    video_client.GetImageSample()

                )

            except Exception as e:

                print(
                    "[카메라 수신 오류]"
                )

                print(e)

                time.sleep(
                    0.2
                )

                continue


            # ---------------------------------------------
            # SDK 오류 코드
            # ---------------------------------------------
            if code != 0:

                time.sleep(
                    0.05
                )

                continue


            # ---------------------------------------------
            # 빈 데이터 방지
            # Go2가 꺼졌거나 영상 수신이 끊긴 경우
            # ---------------------------------------------
            if data is None:

                time.sleep(
                    0.05
                )

                continue


            try:

                if len(data) == 0:

                    time.sleep(
                        0.05
                    )

                    continue

            except TypeError:

                time.sleep(
                    0.05
                )

                continue


            # ---------------------------------------------
            # JPEG 데이터 -> NumPy
            # ---------------------------------------------
            image_data = np.frombuffer(

                bytes(
                    data
                ),

                dtype=np.uint8

            )


            # ---------------------------------------------
            # 빈 NumPy 배열 방지
            # ---------------------------------------------
            if image_data.size == 0:

                time.sleep(
                    0.05
                )

                continue


            # ---------------------------------------------
            # NumPy -> OpenCV 이미지
            # ---------------------------------------------
            try:

                frame = cv2.imdecode(

                    image_data,

                    cv2.IMREAD_COLOR

                )

            except cv2.error as e:

                print(
                    "[카메라 이미지 디코딩 오류]"
                )

                print(e)

                time.sleep(
                    0.05
                )

                continue


            # ---------------------------------------------
            # 디코딩 실패 방지
            # ---------------------------------------------
            if frame is None:

                time.sleep(
                    0.05
                )

                continue


            # ---------------------------------------------
            # 실시간 영상 표시
            # ---------------------------------------------
            cv2.imshow(

                window_name,

                frame

            )


            key = (

                cv2.waitKey(
                    1
                )

                &

                0xFF

            )


            # =================================================
            # a 키
            # 현재 장면 + 얼굴 식별 + GPT + TTS
            # =================================================
            if key == ord(
                "a"
            ):

                print("")
                print(
                    "[a] 현재 카메라 화면을 캡처했습니다."
                )


                analysis_frame = (

                    frame.copy()

                )


                analyze_frame(

                    analysis_frame,

                    openai_client,

                    people,

                    face_model

                )


                print("")
                print(
                    "다시 a 키를 누르면 "
                    "현재 장면과 사람을 분석합니다."
                )


            # =================================================
            # q 또는 ESC
            # =================================================
            elif (

                key == ord(
                    "q"
                )

                or

                key == 27

            ):

                break


            time.sleep(

                0.05

            )


    except KeyboardInterrupt:

        print("")
        print(
            "사용자가 프로그램을 중단했습니다."
        )


    finally:

        cv2.destroyAllWindows()


        try:

            pygame.mixer.quit()

        except Exception:

            pass


    print("")
    print(
        "[3] Go2 AI Vision 종료"
    )


# =========================================================
# 프로그램 시작
# =========================================================
if __name__ == "__main__":

    main()